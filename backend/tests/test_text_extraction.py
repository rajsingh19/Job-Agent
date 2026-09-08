from pathlib import Path
import pytest
from fpdf import FPDF
import docx
from app.services.exceptions import (
    FileCorruptedError,
    TextExtractionRequiredError,
    UnsupportedFileTypeError,
)
from app.services.resume.extractors import (
    DOCXResumeExtractor,
    PDFResumeExtractor,
    get_text_extractor,
)


def create_sample_pdf(file_path: Path, content: str = "Jane Doe\njane@example.com\nSoftware Engineer with Python skills."):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in content.split("\n"):
        pdf.cell(200, 10, text=line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(file_path))


def create_sample_docx(file_path: Path, content: str = "John Doe\njohn@example.com\nBackend Developer with FastAPI and Docker."):
    doc = docx.Document()
    for line in content.split("\n"):
        doc.add_paragraph(line)
    doc.save(str(file_path))


@pytest.mark.asyncio
async def test_pdf_extraction_success(tmp_path: Path):
    pdf_path = tmp_path / "valid_resume.pdf"
    create_sample_pdf(pdf_path, "Jane Candidate\nEmail: jane.c@example.com\nEducation: B.S. CS\nSkills: Python, FastAPI")

    extractor = PDFResumeExtractor()
    extracted_text = await extractor.extract(pdf_path)

    assert "Jane Candidate" in extracted_text
    assert "jane.c@example.com" in extracted_text
    assert "FastAPI" in extracted_text


@pytest.mark.asyncio
async def test_docx_extraction_success(tmp_path: Path):
    docx_path = tmp_path / "valid_resume.docx"
    create_sample_docx(docx_path, "John Applicant\nEmail: john.a@example.com\nSkills: TypeScript, Next.js, PostgreSQL")

    extractor = DOCXResumeExtractor()
    extracted_text = await extractor.extract(docx_path)

    assert "John Applicant" in extracted_text
    assert "john.a@example.com" in extracted_text
    assert "Next.js" in extracted_text


@pytest.mark.asyncio
async def test_empty_or_scanned_pdf_error(tmp_path: Path):
    empty_pdf = tmp_path / "empty_resume.pdf"
    # Create an empty page PDF with no text
    pdf = FPDF()
    pdf.add_page()
    pdf.output(str(empty_pdf))

    extractor = PDFResumeExtractor()
    with pytest.raises(TextExtractionRequiredError):
        await extractor.extract(empty_pdf)


@pytest.mark.asyncio
async def test_corrupted_pdf_error(tmp_path: Path):
    corrupted_pdf = tmp_path / "corrupted.pdf"
    corrupted_pdf.write_bytes(b"This is not a real PDF binary stream.")

    extractor = PDFResumeExtractor()
    with pytest.raises(FileCorruptedError):
        await extractor.extract(corrupted_pdf)


@pytest.mark.asyncio
async def test_corrupted_docx_error(tmp_path: Path):
    corrupted_docx = tmp_path / "corrupted.docx"
    corrupted_docx.write_bytes(b"Not a zip or docx archive.")

    extractor = DOCXResumeExtractor()
    with pytest.raises(FileCorruptedError):
        await extractor.extract(corrupted_docx)


def test_extractor_factory():
    assert isinstance(get_text_extractor("resume.pdf"), PDFResumeExtractor)
    assert isinstance(get_text_extractor("resume.docx"), DOCXResumeExtractor)

    with pytest.raises(UnsupportedFileTypeError):
        get_text_extractor("resume.txt")
