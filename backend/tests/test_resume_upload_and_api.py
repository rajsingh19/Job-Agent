import io
from pathlib import Path
import pytest
from fpdf import FPDF
import docx
from httpx import AsyncClient


def create_pdf_bytes(content: str = "Test User\ntest@example.com\nPython, FastAPI") -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in content.split("\n"):
        pdf.cell(200, 10, text=line, new_x="LMARGIN", new_y="NEXT")
    # Return PDF bytes using bytearray or io
    buffer = io.BytesIO()
    pdf.output(buffer)
    return buffer.getvalue()


def create_docx_bytes(content: str = "Test User\ntest@example.com\nTypeScript, Next.js") -> bytes:
    doc = docx.Document()
    for line in content.split("\n"):
        doc.add_paragraph(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_upload_valid_pdf_and_parse(client: AsyncClient):
    pdf_bytes = create_pdf_bytes("Alice Developer\nalice.dev@example.com\nSkills: Python, FastAPI, Docker")

    # Upload PDF
    files = {"file": ("alice_resume.pdf", pdf_bytes, "application/pdf")}
    headers = {"X-User-Id": "test_user_alice"}
    response = await client.post("/api/v1/resumes", files=files, headers=headers)

    assert response.status_code == 201
    data = response.json()
    assert data["name"] == "alice_resume.pdf"
    assert data["is_default"] is True
    assert "id" in data
    resume_id = data["id"]

    # Trigger Parse
    parse_response = await client.post(f"/api/v1/resumes/{resume_id}/parse", headers=headers)
    assert parse_response.status_code == 200
    parsed_data = parse_response.json()
    assert parsed_data["email"] == "alice.dev@example.com"
    assert parsed_data["low_confidence"] is False

    # Get Profile
    profile_response = await client.get(f"/api/v1/resumes/{resume_id}/profile", headers=headers)
    assert profile_response.status_code == 200
    assert profile_response.json()["email"] == "alice.dev@example.com"


@pytest.mark.asyncio
async def test_upload_valid_docx(client: AsyncClient):
    docx_bytes = create_docx_bytes("Bob Engineer\nbob@example.com\nSkills: Java, Spring Boot")
    files = {"file": ("bob_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    headers = {"X-User-Id": "test_user_bob"}

    response = await client.post("/api/v1/resumes", files=files, headers=headers)
    assert response.status_code == 201
    assert response.json()["name"] == "bob_resume.docx"


@pytest.mark.asyncio
async def test_upload_unsupported_file_extension(client: AsyncClient):
    files = {"file": ("malicious.exe", b"binary content", "application/octet-stream")}
    response = await client.post("/api/v1/resumes", files=files)

    assert response.status_code == 400
    data = response.json()
    assert data["detail"]["error"] == "UNSUPPORTED_FILE_TYPE"


@pytest.mark.asyncio
async def test_upload_oversized_file(client: AsyncClient):
    # 11MB file (exceeds 10MB limit)
    large_bytes = b"0" * (11 * 1024 * 1024)
    files = {"file": ("large_resume.pdf", large_bytes, "application/pdf")}
    response = await client.post("/api/v1/resumes", files=files)

    assert response.status_code == 413
    data = response.json()
    assert data["detail"]["error"] == "FILE_TOO_LARGE"


@pytest.mark.asyncio
async def test_path_traversal_filename_sanitization(client: AsyncClient):
    pdf_bytes = create_pdf_bytes("Charlie Candidate\ncharlie@example.com")
    # Path traversal attack filename
    files = {"file": ("../../../../etc/passwd.pdf", pdf_bytes, "application/pdf")}
    headers = {"X-User-Id": "test_user_charlie"}

    response = await client.post("/api/v1/resumes", files=files, headers=headers)
    assert response.status_code == 201
    # Check that the stored name was sanitized to base name only
    assert response.json()["name"] == "passwd.pdf"


@pytest.mark.asyncio
async def test_list_resumes_endpoint(client: AsyncClient):
    headers = {"X-User-Id": "test_user_multi_resume"}
    pdf1 = create_pdf_bytes("Multi 1\nm1@example.com")
    pdf2 = create_pdf_bytes("Multi 2\nm2@example.com")

    await client.post("/api/v1/resumes", files={"file": ("r1.pdf", pdf1, "application/pdf")}, headers=headers)
    await client.post("/api/v1/resumes", files={"file": ("r2.pdf", pdf2, "application/pdf")}, headers=headers)

    response = await client.get("/api/v1/resumes", headers=headers)
    assert response.status_code == 200
    resumes = response.json()
    assert len(resumes) == 2
