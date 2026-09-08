from abc import ABC, abstractmethod
import os
from pathlib import Path
import pypdf
import docx
from app.services.exceptions import (
    FileCorruptedError,
    TextExtractionFailedError,
    TextExtractionRequiredError,
    UnsupportedFileTypeError,
)


class ResumeTextExtractor(ABC):
    """Abstract interface for resume text extraction."""

    @abstractmethod
    async def extract(self, file_path: Path | str) -> str:
        """Extract clean text content from the file."""
        pass


class PDFResumeExtractor(ResumeTextExtractor):
    """Extracts text content from PDF resume documents."""

    async def extract(self, file_path: Path | str) -> str:
        path = Path(file_path)
        if not path.exists() or not path.is_file():
            raise FileCorruptedError(f"Resume file not found at path: {path.name}")

        try:
            reader = pypdf.PdfReader(str(path))

            if reader.is_encrypted:
                try:
                    # Attempt to decrypt with empty password (common in some PDFs)
                    decrypted = reader.decrypt("")
                    if decrypted == 0:
                        raise FileCorruptedError("PDF is password-protected and cannot be parsed.")
                except Exception:
                    raise FileCorruptedError("PDF is password-protected and cannot be parsed.")

            if len(reader.pages) == 0:
                raise TextExtractionFailedError("PDF document has no pages.")

            extracted_chunks = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    extracted_chunks.append(page_text)

            full_text = "\n\n".join(extracted_chunks).strip()

            # Check if document is scanned or image-only
            alphanumeric_chars = sum(c.isalnum() for c in full_text)
            if alphanumeric_chars < 20:
                raise TextExtractionRequiredError(
                    "PDF document contains insufficient selectable text (scanned or image-only)."
                )

            return full_text

        except (FileCorruptedError, TextExtractionRequiredError, TextExtractionFailedError):
            raise
        except pypdf.errors.PdfReadError as e:
            raise FileCorruptedError(f"Corrupted or invalid PDF file: {str(e)}")
        except Exception as e:
            raise TextExtractionFailedError(f"Unexpected error extracting text from PDF: {str(e)}")


class DOCXResumeExtractor(ResumeTextExtractor):
    """Extracts text content from DOCX resume documents."""

    async def extract(self, file_path: Path | str) -> str:
        path = Path(file_path)
        if not path.exists() or not path.is_file():
            raise FileCorruptedError(f"Resume file not found at path: {path.name}")

        try:
            doc = docx.Document(str(path))
            text_parts = []

            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    text_parts.append(text)

            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        text_parts.append(" | ".join(row_cells))

            full_text = "\n".join(text_parts).strip()

            alphanumeric_chars = sum(c.isalnum() for c in full_text)
            if alphanumeric_chars < 20:
                raise TextExtractionRequiredError(
                    "DOCX document contains insufficient text."
                )

            return full_text

        except (FileCorruptedError, TextExtractionRequiredError, TextExtractionFailedError):
            raise
        except Exception as e:
            raise FileCorruptedError(f"Corrupted or invalid DOCX file: {str(e)}")


def get_text_extractor(filename_or_ext: str) -> ResumeTextExtractor:
    """Factory helper to obtain the appropriate extractor for the file extension."""
    ext = os.path.splitext(filename_or_ext)[1].lower()
    if ext == ".pdf":
        return PDFResumeExtractor()
    elif ext in [".docx", ".doc"]:
        return DOCXResumeExtractor()
    else:
        raise UnsupportedFileTypeError(f"Unsupported extension '{ext}'. Only PDF and DOCX are supported.")
